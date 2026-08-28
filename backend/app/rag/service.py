import io
import structlog
from typing import List, Dict, Any, Optional
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, text
from app.config import settings
from app.db.models import UploadedDocument, DocumentChunk
from app.rag.chunking import split_text_into_chunks
from app.rag.embeddings import get_embedding, get_embeddings_batch, cosine_similarity

logger = structlog.get_logger(__name__)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from PDF stream using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        extracted_text = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_text.append(f"--- Page {i+1} ---\n{page_text}")
        return "\n\n".join(extracted_text)
    except Exception as e:
        logger.error("Error reading PDF file", error=str(e))
        return ""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from DOCX stream using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        extracted_text = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(extracted_text)
    except Exception as e:
        logger.error("Error reading DOCX file", error=str(e))
        return ""


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract plain text from TXT or markdown."""
    try:
        return file_bytes.decode("utf-8", errors="replace")
    except Exception as e:
        logger.error("Error reading text file", error=str(e))
        return ""


async def ingest_document(
    db: AsyncSession,
    student_id: str,
    filename: str,
    original_name: str,
    file_bytes: bytes,
    mime_type: str,
    document_type: str = "general"
) -> UploadedDocument:
    """
    Validate, extract, chunk, embed, and store document chunks in pgvector.
    """
    # 1. Extract text
    content_text = ""
    if "pdf" in mime_type.lower() or filename.lower().endswith(".pdf"):
        content_text = extract_text_from_pdf(file_bytes)
    elif "docx" in mime_type.lower() or "word" in mime_type.lower() or filename.lower().endswith(".docx"):
        content_text = extract_text_from_docx(file_bytes)
    else:
        content_text = extract_text_from_txt(file_bytes)

    if not content_text.strip():
        content_text = f"[Document content from {original_name}]"

    # 2. Split into chunks
    raw_chunks = split_text_into_chunks(
        content_text,
        chunk_size=settings.RAG_CHUNK_SIZE,
        chunk_overlap=settings.RAG_CHUNK_OVERLAP,
        metadata={"filename": original_name, "document_type": document_type}
    )

    # 3. Create document record
    doc = UploadedDocument(
        student_id=student_id,
        filename=filename,
        original_name=original_name,
        mime_type=mime_type,
        file_size=len(file_bytes),
        total_chunks=len(raw_chunks),
        document_type=document_type,
        status="processed"
    )
    db.add(doc)
    await db.flush()

    # 4. Generate embeddings in batch
    texts = [c["content"] for c in raw_chunks]
    embeddings = await get_embeddings_batch(texts)

    # 5. Insert document chunk records
    for i, chunk_data in enumerate(raw_chunks):
        embedding_vec = embeddings[i] if i < len(embeddings) else None
        chunk_obj = DocumentChunk(
            document_id=doc.id,
            chunk_index=chunk_data["chunk_index"],
            content=chunk_data["content"],
            embedding=embedding_vec,
            token_count=chunk_data["token_count"],
            metadata_json=chunk_data["metadata"]
        )
        db.add(chunk_obj)

    await db.commit()
    await db.refresh(doc)
    logger.info("Document ingested successfully", doc_id=doc.id, chunks=len(raw_chunks))
    return doc


async def retrieve_relevant_chunks(
    db: AsyncSession,
    student_id: str,
    query: str,
    top_k: int = settings.RAG_TOP_K,
    similarity_threshold: float = settings.RAG_SIMILARITY_THRESHOLD
) -> List[Dict[str, Any]]:
    """
    Search pgvector for document chunks relevant to the user query for this student.
    Falls back gracefully to text/cosine matching if vector index isn't available.
    """
    query_vector = await get_embedding(query)

    try:
        # Vector search query using pgvector cosine distance operator <=>
        stmt = text("""
            SELECT c.id, c.document_id, c.chunk_index, c.content, c.metadata_json,
                   d.original_name, d.document_type,
                   (1 - (c.embedding <=> :query_embedding::vector)) as similarity
            FROM document_chunks c
            JOIN uploaded_documents d ON c.document_id = d.id
            WHERE d.student_id = :student_id AND c.embedding IS NOT NULL
            ORDER BY c.embedding <=> :query_embedding::vector
            LIMIT :top_k
        """)

        result = await db.execute(stmt, {
            "query_embedding": str(query_vector),
            "student_id": student_id,
            "top_k": top_k
        })
        rows = result.fetchall()

        relevant_chunks = []
        for row in rows:
            sim = float(row.similarity) if row.similarity is not None else 0.0
            if sim >= similarity_threshold:
                relevant_chunks.append({
                    "chunk_id": str(row.id),
                    "document_id": str(row.document_id),
                    "document_name": row.original_name,
                    "document_type": row.document_type,
                    "chunk_index": row.chunk_index,
                    "content": row.content,
                    "similarity": round(sim, 4),
                    "metadata": row.metadata_json or {}
                })

        if relevant_chunks:
            return relevant_chunks

    except Exception as e:
        logger.warning("pgvector SQL search failed or not supported in current connection, using fallback chunk search", error=str(e))

    # Python-level fallback retrieval
    stmt_all = (
        select(DocumentChunk, UploadedDocument)
        .join(UploadedDocument, DocumentChunk.document_id == UploadedDocument.id)
        .where(UploadedDocument.student_id == student_id)
    )
    res = await db.execute(stmt_all)
    all_chunks = res.all()

    scored = []
    for chunk, doc in all_chunks:
        score = 0.5
        if chunk.embedding is not None:
            score = cosine_similarity(query_vector, list(chunk.embedding))
        else:
            # Keyword overlap fallback
            q_words = set(query.lower().split())
            c_words = set(chunk.content.lower().split())
            if q_words:
                overlap = len(q_words.intersection(c_words))
                score = 0.5 + (overlap / len(q_words)) * 0.5

        scored.append((score, chunk, doc))

    scored.sort(key=lambda x: x[0], reverse=True)
    top_results = scored[:top_k]

    return [
        {
            "chunk_id": str(chunk.id),
            "document_id": str(doc.id),
            "document_name": doc.original_name,
            "document_type": doc.document_type,
            "chunk_index": chunk.chunk_index,
            "content": chunk.content,
            "similarity": round(score, 4),
            "metadata": chunk.metadata_json or {}
        }
        for score, chunk, doc in top_results
    ]
